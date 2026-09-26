#!/usr/bin/env python3
"""
Builds Jon Morgan's CV (.docx) from the vault:
  - vault/00-Facts/profile.md              (name, contact, credentials, education)
  - vault/02-Positioning/CV-Positioning.md (headline, summary, quote, competencies)
  - vault/02-Positioning/employment/*.md   (per-employer CV paragraphs)
  - vault/02-Positioning/CV-Project-Highlights.md (per-project blurbs)
  - cv/config.json                          (which projects + employer order to show)

This script only ASSEMBLES text that already exists in the vault. It never
invents wording -- if something reads wrong, fix it in vault/02-Positioning,
not here.

Usage: python build_cv.py [--output ../output/Jon_Morgan_CV.docx]
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent / "lib"))
import load_vault as vault  # noqa: E402

# TODO: confirm this matches the exact "signal red" hex used in the earlier
# CV theme -- this is a reasonable placeholder, not pulled from a saved value.
ACCENT_COLOR = RGBColor(0xC4, 0x12, 0x30)
ACCENT_HEX = "C41230"
TEXT_COLOR = RGBColor(0x1A, 0x1A, 0x1A)
MUTED_COLOR = RGBColor(0x55, 0x55, 0x55)
FONT_NAME = "Calibri"


def add_run(paragraph, text, size_pt, color=None, bold=False, italic=False):
    run = paragraph.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def set_bottom_border(paragraph, color_hex=ACCENT_HEX, size=8):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color_hex)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    set_bottom_border(p)
    add_run(p, text.upper(), size_pt=10, color=ACCENT_COLOR, bold=True)
    return p


def body_paragraph(doc, text, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    add_run(p, text, size_pt=10, color=TEXT_COLOR, italic=italic)
    return p


def bullet_paragraph(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(10)
    run.font.color.rgb = TEXT_COLOR
    return p


def build(vault_dir: Path, config: dict) -> Document:
    profile = vault.load_profile(vault_dir)
    positioning = vault.load_positioning(vault_dir)
    employment = vault.load_employment(vault_dir, config["employmentOrder"])
    highlights = vault.load_project_highlights(vault_dir, config["selectedProjects"])

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(36)
    section.bottom_margin = Pt(36)
    section.left_margin = Pt(45)
    section.right_margin = Pt(45)

    # --- Header: name, headline, contact ---
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_run(p, profile["name"].upper(), size_pt=20, color=ACCENT_COLOR, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_run(p, positioning["headline"], size_pt=11, color=TEXT_COLOR)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    contact = f"{profile['email']}   |   {profile['phone']}   |   {profile['linkedin']}   |   {profile['location']}"
    add_run(p, contact, size_pt=9, color=MUTED_COLOR)

    # --- Summary + quote + target roles ---
    body_paragraph(doc, positioning["summary"])
    quote = positioning["quote"].strip('"“”')
    body_paragraph(doc, f"“{quote}”", italic=True)
    body_paragraph(doc, f"Target roles: {positioning['target_roles']}")

    # --- Core competencies ---
    section_heading(doc, "Core Competencies")
    for c in positioning["core_competencies"]:
        bullet_paragraph(doc, c)

    # --- Professional experience ---
    section_heading(doc, "Professional Experience")
    for job in employment:
        date_loc = ", ".join(x for x in [job.get("period"), job.get("location")] if x)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        add_run(p, job["display_label"], size_pt=10, color=TEXT_COLOR, bold=True)
        if date_loc:
            add_run(p, "\t" + date_loc, size_pt=9, color=MUTED_COLOR)

        if job["type"] == "bullet-list":
            for b in vault.bullets_from(job["body"]):
                bullet_paragraph(doc, b)
        else:
            body_paragraph(doc, job["body"])

    # --- Selected project highlights ---
    section_heading(doc, "Selected Project Highlights")
    for h in highlights:
        bullet_paragraph(doc, h["text"])

    # --- Technical & domain expertise ---
    section_heading(doc, "Technical & Domain Expertise")
    body_paragraph(doc, positioning["domain_expertise"])

    # --- Credentials ---
    if profile["credentials"]:
        section_heading(doc, "Credentials")
        for c in profile["credentials"]:
            bullet_paragraph(doc, c)

    # --- Education ---
    section_heading(doc, "Education")
    for e in profile["education"]:
        bullet_paragraph(doc, e)

    return doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="../output/Jon_Morgan_CV.docx")
    args = parser.parse_args()

    cv_dir = Path(__file__).parent
    vault_dir = cv_dir.parent / "vault"
    config = json.loads((cv_dir / "config.json").read_text(encoding="utf-8"))

    doc = build(vault_dir, config)

    output_path = (cv_dir / args.output).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"Wrote {output_path}")


if __name__ == "__main__":
    main()
